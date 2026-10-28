import csv
from datetime import datetime

from service.db import connect, migrate, add_camera, add_employee, insert_event, \
    open_session, close_session
from reports.daily import write_daily_csv, write_daily_xlsx


def make_db(tmp_path):
    conn = connect(tmp_path / "t.db")
    migrate(conn)
    add_camera(conn, "gate")
    add_employee(conn, "Ayesha", "E1")
    return conn


def seed(conn, emp, out, inn):
    e1 = insert_event(conn, emp, 1, "OUT", out)
    sid = open_session(conn, emp, e1, out, out.date().isoformat())
    e2 = insert_event(conn, emp, 1, "IN", inn)
    close_session(conn, sid, e2, inn)


def test_csv_contents(tmp_path):
    conn = make_db(tmp_path)
    seed(conn, 1, datetime(2026, 10, 26, 13, 0), datetime(2026, 10, 26, 13, 15))
    p = write_daily_csv(conn, "2026-10-26", tmp_path)
    rows = list(csv.DictReader(p.open(encoding="utf-8")))
    assert rows[0]["duration_minutes"] == "15.0"
    assert rows[0]["out_at"].startswith("2026-10-26T13:00")
    assert rows[0]["employee_name"] == "Ayesha"


def test_xlsx_is_valid(tmp_path):
    conn = make_db(tmp_path)
    seed(conn, 1, datetime(2026, 10, 26, 13, 0), datetime(2026, 10, 26, 13, 15))
    p = write_daily_xlsx(conn, "2026-10-26", tmp_path)
    assert p.exists()
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb["Trips"]
    assert ws.max_row >= 2 and ws.cell(1, 1).value == "employee_id"
    assert "Summary" in wb.sheetnames
    s = wb["Summary"]
    assert s.cell(2, 1).value == 1 and s.cell(2, 2).value == 1
    assert s.cell(2, 3).value == 15.0


def test_empty_day_still_writes_headers(tmp_path):
    conn = make_db(tmp_path)
    p = write_daily_csv(conn, "2026-01-01", tmp_path)
    assert p.exists()
    assert len(list(csv.DictReader(p.open(encoding="utf-8")))) == 0
